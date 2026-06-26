# =============================================
# 0. УСТАНОВКА КОДИРОВКИ UTF-8 (безопасно)
# =============================================
import os
import sys

os.environ["PYTHONIOENCODING"] = "utf-8"
os.environ["PYTHONUTF8"] = "1"

# Пытаемся переконфигурировать stdout, если это возможно (не в Streamlit Cloud)
try:
    sys.stdout.reconfigure(encoding='utf-8')
except (AttributeError, ValueError, TypeError):
    pass  # просто игнорируем, если не получилось

# =============================================
# 1. ИМПОРТЫ
# =============================================
import re
import json
from typing import List, Dict, Optional
from io import BytesIO

import streamlit as st
import pandas as pd
from docx import Document
from openpyxl import load_workbook
from openai import OpenAI

# =============================================
# 2. ПОЛУЧЕНИЕ API-КЛЮЧА
# =============================================
CLOUD_RU_API_KEY = os.getenv("CLOUD_RU_API_KEY")  # для Hugging Face и Streamlit Cloud Secrets

# =============================================
# 3. АВТООПРЕДЕЛЕНИЕ КОЛОНОК
# =============================================
def detect_text_column(headers: List[str], df_sample: pd.DataFrame) -> Optional[str]:
    keywords = [
        "propertyvalue", "описание", "жалобы", "анамнез", "заключение",
        "текст", "anamnesis", "complaints", "diagnosis", "text",
        "описание случая", "жалобы при поступлении", "жалобы пациента",
        "анамнез заболевания", "notes", "history", "description",
        "симптомы", "диагноз", "пациент", "жалобы", "медицинские записи"
    ]
    for col in headers:
        col_lower = col.lower().strip()
        if any(kw in col_lower for kw in keywords):
            return col
    if df_sample is not None and not df_sample.empty:
        avg_lengths = {}
        for col in df_sample.columns:
            sample = df_sample[col].dropna()
            if len(sample) == 0:
                continue
            try:
                numeric_ratio = pd.to_numeric(sample, errors='coerce').notna().mean()
                if numeric_ratio > 0.5:
                    continue
            except:
                pass
            avg_len = sample.astype(str).str.len().mean()
            avg_lengths[col] = avg_len
        if avg_lengths:
            return max(avg_lengths, key=avg_lengths.get)
    for col in headers:
        if df_sample is not None and df_sample[col].dtype == 'object':
            return col
    return None

def detect_id_column(headers: List[str], df_sample: pd.DataFrame, text_col: Optional[str] = None) -> Optional[str]:
    keywords = ["id", "номер", "код", "patient", "propertyid", "case", "идентификатор"]
    for col in headers:
        col_lower = col.lower().strip()
        if any(kw in col_lower for kw in keywords):
            return col
    if df_sample is not None and not df_sample.empty:
        unique_counts = {}
        for col in df_sample.columns:
            if col == text_col:
                continue
            n_unique = df_sample[col].nunique()
            if n_unique > 2 and n_unique / len(df_sample) > 0.05:
                unique_counts[col] = n_unique
        if unique_counts:
            return max(unique_counts, key=unique_counts.get)
    return None

# =============================================
# 4. ФУНКЦИИ ЧТЕНИЯ ФАЙЛОВ
# =============================================
def read_docx(file_path: str, text_col: Optional[str] = None, id_col: Optional[str] = None) -> List[Dict]:
    doc = Document(file_path)
    records = []
    row_counter = 0
    for table in doc.tables:
        if not table.rows:
            continue
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        if text_col is None:
            sample_rows = []
            for row in table.rows[1:11]:
                sample_rows.append([cell.text.strip() for cell in row.cells])
            if headers and sample_rows:
                df_sample = pd.DataFrame(sample_rows, columns=headers)
                text_col = detect_text_column(headers, df_sample)
                id_col = detect_id_column(headers, df_sample, text_col)
            else:
                continue
        if text_col is None or text_col not in headers:
            continue
        text_idx = headers.index(text_col)
        id_idx = headers.index(id_col) if id_col and id_col in headers else None
        for row in table.rows[1:]:
            row_counter += 1
            cells = row.cells
            if text_idx < len(cells):
                text = cells[text_idx].text.strip()
                if text:
                    patient_id = None
                    if id_idx is not None and id_idx < len(cells):
                        patient_id = cells[id_idx].text.strip() or None
                    records.append({
                        'row_num': row_counter,
                        'text': text,
                        'source': f"{os.path.basename(file_path)}, строка {row_counter}",
                        'patient_id': patient_id
                    })
    return records

def read_xlsx(file_path: str, text_col: Optional[str] = None, id_col: Optional[str] = None) -> List[Dict]:
    wb = load_workbook(file_path, data_only=True)
    records = []
    row_counter = 0
    for sheet in wb.worksheets:
        data = sheet.values
        headers = None
        rows_data = []
        for i, row in enumerate(data):
            if i == 0:
                headers = [str(cell) if cell else "" for cell in row]
            else:
                rows_data.append(row)
        if not headers:
            continue
        if text_col is None:
            df_sample = pd.DataFrame(rows_data[:100], columns=headers)
            text_col = detect_text_column(headers, df_sample)
            id_col = detect_id_column(headers, df_sample, text_col)
        if text_col is None or text_col not in headers:
            continue
        text_idx = headers.index(text_col)
        id_idx = headers.index(id_col) if id_col and id_col in headers else None
        for row in rows_data:
            row_counter += 1
            if len(row) > text_idx:
                text = str(row[text_idx]) if row[text_idx] is not None else ""
                if text.strip():
                    patient_id = None
                    if id_idx is not None and len(row) > id_idx and row[id_idx] is not None:
                        patient_id = str(row[id_idx])
                    records.append({
                        'row_num': row_counter,
                        'text': text.strip(),
                        'source': f"{os.path.basename(file_path)}, лист {sheet.title}, строка {row_counter}",
                        'patient_id': patient_id
                    })
    return records

def read_txt(file_path: str) -> List[Dict]:
    with open(file_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    records = []
    for i, line in enumerate(lines, 1):
        text = line.strip()
        if text:
            records.append({
                'row_num': i,
                'text': text,
                'source': f"{os.path.basename(file_path)}, строка {i}",
                'patient_id': None
            })
    return records

def detect_file_type(file_path: str) -> Optional[str]:
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.docx': return 'docx'
    elif ext == '.xlsx': return 'xlsx'
    elif ext == '.txt': return 'txt'
    return None

def load_folder(folder_path: str, text_col: Optional[str] = None, id_col: Optional[str] = None) -> List[Dict]:
    os.makedirs(folder_path, exist_ok=True)
    all_records = []
    for filename in os.listdir(folder_path):
        file_path = os.path.join(folder_path, filename)
        ft = detect_file_type(file_path)
        if ft == 'docx':
            all_records.extend(read_docx(file_path, text_col, id_col))
        elif ft == 'xlsx':
            all_records.extend(read_xlsx(file_path, text_col, id_col))
        elif ft == 'txt':
            all_records.extend(read_txt(file_path))
    return all_records

# =============================================
# 5. АДАПТЕР ДЛЯ CLOUD.RU
# =============================================
class CloudRuAdapter:
    def __init__(self, model: str = "Qwen/Qwen3-30B-A3B", api_key: str = None, timeout: int = 300):
        self.model = model
        self.api_key = api_key
        if not self.api_key:
            raise ValueError("API-ключ Cloud.ru не указан")
        self.client = OpenAI(
            base_url="https://foundation-models.api.cloud.ru/v1",
            api_key=self.api_key,
            timeout=timeout
        )

    def generate(self, system_prompt: str, user_prompt: str, **kwargs) -> str:
        messages = []
        if system_prompt:
            messages.append({"role": "system", "content": system_prompt})
        messages.append({"role": "user", "content": user_prompt})
        response = self.client.chat.completions.create(
            model=self.model,
            messages=messages,
            temperature=kwargs.get("temperature", 0.5),
            max_tokens=kwargs.get("max_tokens", 2500),
            presence_penalty=kwargs.get("presence_penalty", 0),
            top_p=kwargs.get("top_p", 0.95),
        )
        return response.choices[0].message.content

# =============================================
# 6. ИЗВЛЕЧЕНИЕ ПАР (с защитой от ошибок)
# =============================================
def extract_pairs(text: str, llm) -> list:
    if len(text) > 1500:
        text = text[:1500] + "..."
    prompt = f"""
Извлеки из текста все медицинские термины (симптомы, диагнозы, состояния) и их наличие/значение.
Верни JSON-массив массивов, каждый подмассив из двух строк: ["термин", "значение"].
Возможные значения: присутствует, отсутствует, усиливается, ослабевает, постоянный, периодический и т.п.
Только JSON, без пояснений.

Текст:
{text}
"""
    content = llm.generate(
        system_prompt="Ты медицинский эксперт. Отвечай только JSON.",
        user_prompt=prompt,
        temperature=0
    )
    content = content.strip().removeprefix("```json").removesuffix("```").strip()
    try:
        data = json.loads(content)
    except:
        match = re.search(r'\[.*\]', content, re.DOTALL)
        if match:
            try:
                data = json.loads(match.group())
            except:
                data = []
        else:
            data = []
    if isinstance(data, dict) and "terms" in data:
        data = data["terms"]
    if not isinstance(data, list):
        data = []
    # Принудительно приводим к списку из двух элементов
    filtered = []
    for item in data:
        if isinstance(item, (list, tuple)) and len(item) >= 2:
            filtered.append([item[0], item[1]])
    return filtered

# =============================================
# 7. ДОБАВЛЕНИЕ ХАРАКТЕРИСТИК
# =============================================
def enrich_with_characteristics(pairs: list, llm) -> list:
    unique_values = set(val for _, val in pairs)
    value_to_char = {}

    for value in unique_values:
        prompt = f"""
Определи категорию (характеристику) для медицинского значения "{value}".
Категория должна быть существительным или короткой фразой, обобщающей тип значения.
Примеры:
- "присутствует" → "присутствие"
- "отсутствует" → "отсутствие"
- "усиливается" → "динамика"
- "постоянный" → "характер"
- "ноющая" → "интенсивность" (или "характер боли")
- "при нагрузке" → "условие возникновения"
- "5 мг" → "дозировка"

Верни ТОЛЬКО категорию, без пояснений. Одно слово или короткая фраза.
"""
        try:
            response = llm.generate(
                system_prompt="Ты эксперт по медицинской терминологии.",
                user_prompt=prompt,
                temperature=0
            )
            char = response.strip().strip('"').strip("'")
            if not char or len(char) > 100:
                char = "характеристика"
            value_to_char[value] = char
        except Exception as e:
            value_to_char[value] = "характеристика"

    triples = []
    for term, value in pairs:
        char = value_to_char.get(value, "характеристика")
        triples.append((term, char, value))
    return triples

# =============================================
# 8. ОСНОВНОЙ ПАЙПЛАЙН
# =============================================
def run_pipeline(folder: str, model: str, api_key: str,
                 text_col: Optional[str] = None, id_col: Optional[str] = None,
                 timeout: int = 300):
    llm = CloudRuAdapter(api_key=api_key, model=model, timeout=timeout)
    os.makedirs(folder, exist_ok=True)

    data = load_folder(folder_path=folder, text_col=text_col, id_col=id_col)
    if not data:
        return None, None, None

    all_pairs = []
    for rec in data:
        pairs = extract_pairs(rec['text'], llm)
        if not pairs:
            continue
        for item in pairs:
            t, v = item[0], item[1]
            all_pairs.append({
                'term': t,
                'value': v,
                'source': rec['source'],
                'source_text': rec['text'],
                'patient_id': rec.get('patient_id')
            })
    if not all_pairs:
        return data, [], []

    pair_list = [(p['term'], p['value']) for p in all_pairs]
    triples = enrich_with_characteristics(pair_list, llm)

    enriched = [{
        'term': t,
        'characteristic': c,
        'value': v,
        'source': p['source'],
        'source_text': p['source_text'],
        'patient_id': p['patient_id']
    } for (t, c, v), p in zip(triples, all_pairs)]

    return data, all_pairs, enriched

# =============================================
# 9. STREAMLIT ИНТЕРФЕЙС
# =============================================
st.set_page_config(page_title="Медицинский парсер", layout="centered")
st.title("🏥 Преобразование медицинских записей в датасет")
st.markdown("Загрузите файлы (Excel, Word, TXT). Колонки определяются автоматически, но вы можете уточнить.")

# API-ключ
if CLOUD_RU_API_KEY:
    api_key_input = CLOUD_RU_API_KEY
    st.success("🔑 API-ключ загружен из секретов")
else:
    api_key_input = st.text_input("Введите ваш API-ключ Cloud.ru", type="password")

# Загрузка файлов
uploaded_files = st.file_uploader(
    "Выберите файлы",
    accept_multiple_files=True,
    type=['xlsx', 'docx', 'txt']
)

text_col = None
id_col = None

if uploaded_files:
    xlsx_files = [f for f in uploaded_files if f.name.endswith('.xlsx')]
    if xlsx_files:
        try:
            df_sample = pd.read_excel(xlsx_files[0], nrows=1)
            headers = df_sample.columns.tolist()
            if headers:
                temp_text = detect_text_column(headers, df_sample)
                temp_id = detect_id_column(headers, df_sample, temp_text)
                st.write("Выберите колонки для обработки (или оставьте автоопределение):")
                col1, col2 = st.columns(2)
                with col1:
                    text_col = st.selectbox(
                        "Колонка с текстом (описаниями):",
                        options=["(автоопределение)"] + headers,
                        index=0 if temp_text is None else headers.index(temp_text) + 1
                    )
                    if text_col == "(автоопределение)":
                        text_col = None
                with col2:
                    id_col = st.selectbox(
                        "Колонка с ID пациента (если есть):",
                        options=["(автоопределение)"] + headers,
                        index=0 if temp_id is None else headers.index(temp_id) + 1
                    )
                    if id_col == "(автоопределение)":
                        id_col = None
        except Exception as e:
            st.warning(f"Не удалось прочитать заголовки: {e}")

if st.button("Обработать"):
    if not api_key_input:
        st.error("Пожалуйста, введите API-ключ.")
    elif not uploaded_files:
        st.error("Загрузите хотя бы один файл.")
    else:
        input_dir = "input"
        os.makedirs(input_dir, exist_ok=True)
        for f in os.listdir(input_dir):
            os.remove(os.path.join(input_dir, f))
        for uploaded_file in uploaded_files:
            with open(os.path.join(input_dir, uploaded_file.name), "wb") as f:
                f.write(uploaded_file.getbuffer())

        with st.spinner("Идёт обработка... Это может занять несколько минут."):
            try:
                data, all_pairs, enriched = run_pipeline(
                    folder="input",
                    model="Qwen/Qwen3-30B-A3B",
                    api_key=api_key_input,
                    text_col=text_col,
                    id_col=id_col
                )

                if data is None:
                    st.error("Не удалось загрузить данные. Проверьте формат файлов.")
                elif not all_pairs:
                    st.warning("Не найдено ни одной пары (термин-значение). Возможно, текст не содержит медицинских терминов.")
                else:
                    output = BytesIO()
                    with pd.ExcelWriter(output, engine='openpyxl') as writer:
                        pd.DataFrame(data).to_excel(writer, sheet_name='Исходные записи', index=False)
                        pd.DataFrame(all_pairs).to_excel(writer, sheet_name='Пары', index=False)
                        pd.DataFrame(enriched).to_excel(writer, sheet_name='Итог', index=False)
                        if enriched:
                            char_dict = {r['value']: r['characteristic'] for r in enriched}
                            pd.DataFrame(list(char_dict.items()), columns=['Значение', 'Характеристика']).to_excel(writer, sheet_name='Словарь характеристик', index=False)
                    output.seek(0)

                    st.success(f"Обработка завершена! Извлечено {len(enriched)} троек.")
                    st.download_button(
                        label="📥 Скачать result.xlsx",
                        data=output,
                        file_name="result.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
                    json_output = BytesIO()
                    json_output.write(json.dumps(enriched, ensure_ascii=False, indent=2).encode('utf-8'))
                    json_output.seek(0)
                    st.download_button(
                        label="📥 Скачать result.json",
                        data=json_output,
                        file_name="result.json",
                        mime="application/json"
                    )
            except Exception as e:
                st.error(f"Ошибка: {e}")
                st.code(str(e))

if __name__ == "__main__":
    pass
